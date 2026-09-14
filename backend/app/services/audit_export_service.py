"""
Модуль для генерации отчетов по журналу аудита системы в формате Excel (.xlsx).
Предоставляет функции преобразования технических логов аудита в человекочитаемые описания
и экспорта структурированного отчета с корпоративным оформлением.
"""
import io
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.models.user import User
from app.core.config import settings


# Словарь человекопонятных названий действий на русском языке
ACTION_TITLES: dict[str, str] = {
    "LOGIN": "Вход в систему",
    "LOGOUT": "Выход из системы",
    "PASSWORD_RESET_REQUEST": "Запрос сброса пароля",
    "PASSWORD_RESET_CONFIRM": "Подтверждение сброса пароля",
    "CHANGE_PASSWORD": "Смена пароля",
    "RESET_PASSWORD": "Сброс пароля администратором",
    "CREATE_USER": "Создание пользователя",
    "UPDATE_USER": "Обновление данных пользователя",
    "DELETE_USER": "Удаление пользователя",
    "CREATE_DEPT": "Создание отдела",
    "UPDATE_DEPT": "Изменение отдела",
    "DELETE_DEPT": "Удаление отдела",
    "CREATE_PROJECT": "Создание проекта",
    "UPDATE_PROJECT": "Изменение проекта",
    "DELETE_PROJECT": "Удаление проекта",
    "IMPORT_ODOO_PROJECTS": "Импорт проектов (Odoo XML-RPC)",
    "IMPORT_ODOO_INTEGRATION_PROJECTS": "Импорт проектов (Odoo API)",
    "CREATE_OVERTIME": "Создание заявки на переработку",
    "UPDATE_OVERTIME_TIME": "Изменение времени переработки",
    "REVIEW_HEAD_APPROVED": "Согласование руководителем (Одобрено)",
    "REVIEW_HEAD_REJECTED": "Согласование руководителем (Отклонено)",
    "REVIEW_MANAGER_APPROVED": "Согласование менеджером (Одобрено)",
    "REVIEW_MANAGER_REJECTED": "Согласование менеджером (Отклонено)",
    "CANCEL_OVERTIME": "Отмена заявки",
    "RESTORE_OVERTIME": "Восстановление заявки",
    "AUTO_CLOSE_STALE": "Автоматическое закрытие заявки",
}

# Категории действий для группировки
ACTION_CATEGORIES: dict[str, str] = {
    "LOGIN": "Авторизация",
    "LOGOUT": "Авторизация",
    "PASSWORD_RESET_REQUEST": "Безопасность",
    "PASSWORD_RESET_CONFIRM": "Безопасность",
    "CHANGE_PASSWORD": "Безопасность",
    "RESET_PASSWORD": "Безопасность",
    "CREATE_USER": "Пользователи",
    "UPDATE_USER": "Пользователи",
    "DELETE_USER": "Пользователи",
    "CREATE_DEPT": "Отделы",
    "UPDATE_DEPT": "Отделы",
    "DELETE_DEPT": "Отделы",
    "CREATE_PROJECT": "Проекты",
    "UPDATE_PROJECT": "Проекты",
    "DELETE_PROJECT": "Проекты",
    "IMPORT_ODOO_PROJECTS": "Интеграция Odoo",
    "IMPORT_ODOO_INTEGRATION_PROJECTS": "Интеграция Odoo",
    "CREATE_OVERTIME": "Заявки",
    "UPDATE_OVERTIME_TIME": "Заявки",
    "REVIEW_HEAD_APPROVED": "Согласование",
    "REVIEW_HEAD_REJECTED": "Согласование",
    "REVIEW_MANAGER_APPROVED": "Согласование",
    "REVIEW_MANAGER_REJECTED": "Согласование",
    "CANCEL_OVERTIME": "Заявки",
    "RESTORE_OVERTIME": "Заявки",
    "AUTO_CLOSE_STALE": "Система",
}


def format_dt_local(dt: datetime | None) -> str:
    """
    Преобразует datetime объект в строку локального времени организации.
    Формат: ДД.ММ.ГГГГ ЧЧ:ММ:СС.
    """
    if not dt:
        return "-"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    local_dt = dt.astimezone(settings.tz_info)
    return local_dt.strftime("%d.%m.%Y %H:%M:%S")


def format_details_to_text(action: str, details: dict | None) -> str:
    """
    Формирует понятное человекочитаемое текстовое описание деталей события аудита.
    """
    if not details or not isinstance(details, dict):
        return "-"

    parts: list[str] = []

    if action == "UPDATE_OVERTIME_TIME":
        old_h = details.get("old_hours")
        new_h = details.get("new_hours")
        old_s = details.get("old_start")
        new_s = details.get("new_start")
        old_e = details.get("old_end")
        new_e = details.get("new_end")
        updated_by = details.get("updated_by")
        role = details.get("role")

        if updated_by:
            parts.append(f"Инициатор: {updated_by} ({role or 'роль не указана'})")
        if old_h is not None and new_h is not None:
            parts.append(f"Часы: было {old_h} ч. → стало {new_h} ч.")
        if old_s and new_s:
            parts.append(f"Начало: {old_s[:16]} → {new_s[:16]}")
        if old_e and new_e:
            parts.append(f"Окончание: {old_e[:16]} → {new_e[:16]}")

    elif action.startswith("REVIEW_"):
        approved = details.get("approved")
        app_h = details.get("approved_hours")
        req_h = details.get("requested_hours")
        comment = details.get("comment")

        parts.append("Решение: Одобрено" if approved else "Решение: Отклонено")
        if app_h is not None:
            parts.append(f"Одобрено часов: {app_h} ч. (запрошено: {req_h} ч.)")
        if comment:
            parts.append(f"Комментарий: «{comment}»")

    elif action == "CANCEL_OVERTIME":
        by = details.get("cancelled_by")
        prev = details.get("previous_status")
        desc = details.get("description")
        if by:
            parts.append(f"Отменил: {by}")
        if prev:
            parts.append(f"Предыдущий статус: {prev}")
        if desc:
            parts.append(f"Описание: «{desc}»")

    elif action == "AUTO_CLOSE_STALE":
        reason = details.get("reason")
        auto_end = details.get("auto_end")
        if reason:
            parts.append(f"Причина: {reason}")
        if auto_end:
            parts.append(f"Авто-завершение: {auto_end}")

    elif action in ("CREATE_USER", "UPDATE_USER"):
        name = details.get("full_name")
        email = details.get("email")
        role = details.get("role")
        is_active = details.get("is_active")
        if name:
            parts.append(f"ФИО: {name}")
        if email:
            parts.append(f"Email: {email}")
        if role:
            parts.append(f"Роль: {role}")
        if is_active is not None:
            parts.append(f"Активен: {'Да' if is_active else 'Нет'}")

    elif action in ("CREATE_DEPT", "UPDATE_DEPT"):
        name = details.get("name")
        head_id = details.get("head_id")
        if name:
            parts.append(f"Название: {name}")
        if "head_id" in details:
            parts.append(f"ID руководителя: {head_id if head_id is not None else 'Не назначен'}")

    elif action in ("CREATE_PROJECT", "UPDATE_PROJECT"):
        name = details.get("name")
        code = details.get("code")
        limit = details.get("weekly_limit")
        is_active = details.get("is_active")
        if name:
            parts.append(f"Название: {name}")
        if code:
            parts.append(f"Код: {code}")
        if limit is not None:
            parts.append(f"Лимит: {limit} ч/нед")
        if is_active is not None:
            parts.append(f"Статус: {'Активен' if is_active else 'В архиве'}")

    elif action in ("IMPORT_ODOO_PROJECTS", "IMPORT_ODOO_INTEGRATION_PROJECTS"):
        imported = details.get("imported")
        skipped = details.get("skipped")
        if imported is not None:
            parts.append(f"Импортировано: {imported}")
        if skipped is not None:
            parts.append(f"Пропущено: {skipped}")

    else:
        for k, v in details.items():
            if v is not None and str(v).strip():
                parts.append(f"{k}: {v}")

    return "; ".join(parts) if parts else "-"


async def generate_audit_excel_file(
    items: list[dict],
    current_user: User,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
) -> io.BytesIO:
    """
    Генерирует форматированный Excel-файл журнала аудита событий системы.

    :param items: Список словарей записей аудита.
    :param current_user: Администратор, запросивший выгрузку.
    :param start_date: Начало периода отчета.
    :param end_date: Окончание периода отчета.
    :return: BytesIO объект с готовым .xlsx файлом.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Журнал аудита"

    # Стили
    font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
    font_subtitle = Font(name="Arial", size=10, italic=True, color="64748B")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Arial", size=9, color="0F172A")
    font_bold = Font(name="Arial", size=9, bold=True, color="0F172A")

    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Заголовок отчета
    ws.merge_cells("A1:H1")
    ws["A1"] = "ЖУРНАЛ АУДИТА И ИСТОРИИ ДЕЙСТВИЙ В СИСТЕМЕ"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    # 2. Метаданные (период, дата выгрузки, автор)
    period_str = "Период: За все время"
    if start_date and end_date:
        period_str = f"Период: с {format_dt_local(start_date)[:10]} по {format_dt_local(end_date)[:10]}"
    elif start_date:
        period_str = f"Период: с {format_dt_local(start_date)[:10]}"
    elif end_date:
        period_str = f"Период: по {format_dt_local(end_date)[:10]}"

    export_now_str = format_dt_local(datetime.now(timezone.utc))

    ws["A2"] = f"{period_str}  |  Дата выгрузки: {export_now_str}  |  Выгрузил: {current_user.full_name}"
    ws["A2"].font = font_subtitle
    ws.row_dimensions[2].height = 18

    # 3. Заголовки колонок
    headers = [
        "№",
        "Дата и время",
        "Пользователь",
        "Email",
        "Категория",
        "Действие",
        "Объект",
        "Подробности / Описание изменений"
    ]

    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 28

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    # 4. Заполнение данными
    row_idx = 5
    for idx, item in enumerate(items, start=1):
        raw_action = item.get("action", "")
        action_title = ACTION_TITLES.get(raw_action, raw_action)
        category = ACTION_CATEGORIES.get(raw_action, "Общее")
        
        user_info = item.get("user") or {}
        user_name = user_info.get("full_name") or ("Система" if not item.get("user_id") else f"ID: {item.get('user_id')}")
        user_email = user_info.get("email") or "-"
        
        target_type = item.get("target_type") or "-"
        target_id = item.get("target_id")
        target_str = f"{target_type} #{target_id}" if target_id else target_type

        ts = item.get("timestamp")
        dt_str = format_dt_local(ts) if isinstance(ts, datetime) else str(ts or "-")

        details_desc = format_details_to_text(raw_action, item.get("details"))

        row_data = [
            idx,
            dt_str,
            user_name,
            user_email,
            category,
            action_title,
            target_str,
            details_desc
        ]

        ws.row_dimensions[row_idx].height = 24
        is_even = (idx % 2 == 0)

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_bold if col_idx == 1 else font_body
            cell.border = border_cell
            if is_even:
                cell.fill = fill_zebra

            if col_idx in (1, 2, 5):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

        row_idx += 1

    # 5. Настройка ширины колонок
    col_widths = {
        1: 6,   # №
        2: 20,  # Дата
        3: 26,  # Пользователь
        4: 24,  # Email
        5: 16,  # Категория
        6: 34,  # Действие
        7: 18,  # Объект
        8: 50,  # Подробности
    }

    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Закрепляем область под заголовками
    ws.freeze_panes = "A5"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
