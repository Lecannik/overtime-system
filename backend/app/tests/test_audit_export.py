import pytest
from datetime import datetime, timezone
import openpyxl
from app.services.audit_export_service import (
    generate_audit_excel_file,
    format_details_to_text,
    ACTION_TITLES,
    ACTION_CATEGORIES
)
from app.models.user import User, UserRole, UserCompany


def test_action_titles_and_formatting():
    """
    Проверяет корректность перевода действий и форматирования деталей.
    """
    assert ACTION_TITLES["UPDATE_OVERTIME_TIME"] == "Изменение времени переработки"
    assert ACTION_CATEGORIES["UPDATE_OVERTIME_TIME"] == "Заявки"


    details_time = {
        "updated_by": "Иван Иванов",
        "role": "admin",
        "old_hours": 3.0,
        "new_hours": 4.5,
        "old_start": "2026-08-28T18:00:00Z",
        "new_start": "2026-08-28T17:30:00Z",
        "old_end": "2026-08-28T21:00:00Z",
        "new_end": "2026-08-28T22:00:00Z"
    }
    summary_time = format_details_to_text("UPDATE_OVERTIME_TIME", details_time)
    assert "Часы: было 3.0 ч. → стало 4.5 ч." in summary_time
    assert "Инициатор: Иван Иванов (admin)" in summary_time

    details_rev = {
        "approved": True,
        "approved_hours": 4.0,
        "requested_hours": 4.0,
        "comment": "Согласовано без замечаний"
    }
    summary_rev = format_details_to_text("REVIEW_HEAD_APPROVED", details_rev)
    assert "Решение: Одобрено" in summary_rev
    assert "Одобрено часов: 4.0 ч." in summary_rev
    assert "Комментарий: «Согласовано без замечаний»" in summary_rev




@pytest.mark.asyncio
async def test_generate_audit_excel_file():
    """
    Проверяет создание Excel-файла журнала аудита.
    """
    admin_user = User(
        id=1,
        full_name="Администратор Системы",
        email="admin@example.com",
        role=UserRole.admin,
        company=UserCompany.Polymedia,
        is_active=True
    )

    items = [
        {
            "id": 1,
            "action": "UPDATE_OVERTIME_TIME",
            "target_type": "overtime",
            "target_id": 101,
            "user_id": 1,
            "user": {
                "id": 1,
                "full_name": "Администратор Системы",
                "email": "admin@example.com"
            },
            "timestamp": datetime(2026, 8, 28, 10, 0, 0, tzinfo=timezone.utc),
            "details": {
                "updated_by": "Администратор Системы",
                "role": "admin",
                "old_hours": 2.0,
                "new_hours": 3.0
            }
        },
        {
            "id": 2,
            "action": "AUTO_CLOSE_STALE",
            "target_type": "overtime",
            "target_id": 102,
            "user_id": None,
            "user": None,
            "timestamp": datetime(2026, 8, 28, 11, 0, 0, tzinfo=timezone.utc),
            "details": {
                "reason": "Автоматическое закрытие незавершенной переработки"
            }
        }
    ]

    excel_buf = await generate_audit_excel_file(
        items=items,
        current_user=admin_user,
        start_date=datetime(2026, 8, 1, 0, 0, 0, tzinfo=timezone.utc),
        end_date=datetime(2026, 8, 31, 23, 59, 59, tzinfo=timezone.utc)
    )

    excel_bytes = excel_buf.getvalue()
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(excel_buf)
    ws = wb.active
    assert ws.title == "Журнал аудита"
    assert "ЖУРНАЛ АУДИТА" in ws["A1"].value
    assert ws.cell(row=5, column=3).value == "Администратор Системы"
    assert ws.cell(row=5, column=4).value == "admin@example.com"
    assert ws.cell(row=5, column=6).value == "Изменение времени переработки"
    assert ws.cell(row=6, column=3).value == "Система"
